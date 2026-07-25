"""Extract weekly plan/fact tables from stroyka PPTX (embedded Excel OLE)."""

from __future__ import annotations

import io
import re
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass
class WeeklyReport:
    """One work package from a PPT slide (OLE sheet)."""

    slide_index: int
    title: str
    title_ole: str
    period: str
    weeks_plan: list[float | None]
    weeks_fact: list[float | None]
    total: float | None
    done: float | None
    rest: float | None
    unit: str
    pct: float | None
    history_months: dict[str, float] = field(default_factory=dict)
    project_id: str = ""
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _clean_title(t: str) -> str:
    return re.sub(r"\s+", " ", (t or "").replace("\n", " ")).strip()


def _parse_sheet(ws, slide_index: int) -> WeeklyReport:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    # Pad short rows
    width = max((len(r) for r in rows), default=0)
    rows = [r + [None] * (width - len(r)) for r in rows]

    period = ""
    title_ole = ""
    if rows:
        period = str(rows[0][4] or "")
    if len(rows) > 2:
        title_ole = _clean_title(str(rows[2][3] or ""))

    weeks_plan: list[float | None] = []
    weeks_fact: list[float | None] = []
    total = done = rest = pct = None
    unit = ""
    history: dict[str, float] = {}

    # Row layout from contractor template:
    # r2: history months in cols 0..2, title col3, 'план' col4, weeks 5..9, month total 10
    # r3: факт + small table Всего at col 12..16
    # r4: Выполнено
    # r5: Остаток
    if len(rows) > 2:
        r = rows[2]
        for i, label in enumerate(("2026-04", "2026-05", "2026-06")):
            n = _num(r[i]) if i < len(r) else None
            if n is not None:
                history[label] = n
        weeks_plan = [_num(r[i]) for i in range(5, 10)]
    if len(rows) > 3:
        r = rows[3]
        weeks_fact = [_num(r[i]) for i in range(5, 10)]
        # Small table: label, value, unit, pct — у строки «Всего» pct=100 (доля плана), не % выполнения
        for i, cell in enumerate(r):
            if str(cell).strip() == "Всего" and i + 1 < len(r):
                total = _num(r[i + 1])
                unit = str(r[i + 2] or "").strip()
                break
    if len(rows) > 4:
        r = rows[4]
        for i, cell in enumerate(r):
            if str(cell).strip() == "Выполнено" and i + 1 < len(r):
                done = _num(r[i + 1])
                pct = _num(r[i + 3]) if i + 3 < len(r) else None
                break
    if len(rows) > 5:
        r = rows[5]
        for i, cell in enumerate(r):
            if str(cell).strip() == "Остаток" and i + 1 < len(r):
                rest = _num(r[i + 1])
                break

    if total is not None and done is not None and rest is None:
        rest = total - done
    if total and done is not None and total > 0:
        # Всегда считаем % от объёмов (надёжнее ячейки шаблона)
        pct = done / total * 100

    notes: list[str] = []
    unit_norm = unit
    if unit in {"м4", "м5"}:
        notes.append(f"Ед.изм в OLE «{unit}» — нормализуем в м3")
        unit_norm = "м3"

    return WeeklyReport(
        slide_index=slide_index,
        title=title_ole,
        title_ole=title_ole,
        period=period,
        weeks_plan=weeks_plan,
        weeks_fact=weeks_fact,
        total=total,
        done=done,
        rest=rest,
        unit=unit_norm or unit,
        pct=pct,
        history_months=history,
        notes=notes,
    )


def _slide_titles_from_pptx(pptx_bytes: bytes) -> list[str]:
    """Best-effort work titles from slide shapes (skip project id / site name boxes)."""
    try:
        from pptx import Presentation
    except ImportError:
        return []
    uuid_re = re.compile(
        r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}",
        re.I,
    )
    skip_words = ("ленинский", "id проекта", "id проекта")
    prs = Presentation(io.BytesIO(pptx_bytes))
    titles: list[str] = []
    for slide in prs.slides:
        candidates: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = _clean_title("\n".join(p.text for p in shape.text_frame.paragraphs))
            if not t or uuid_re.search(t):
                continue
            low = t.lower()
            if low.startswith("id") or any(w in low for w in skip_words):
                continue
            if t.isdigit():
                continue
            candidates.append(t)
        # Prefer longest meaningful work title
        best = max(candidates, key=len) if candidates else ""
        titles.append(best)
    return titles


def _project_id_from_name(name: str) -> str:
    m = re.search(
        r"([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})",
        name,
        re.I,
    )
    return m.group(1).lower() if m else ""


def parse_stroyka_pptx(path_or_bytes: str | Path | bytes, filename: str = "") -> list[WeeklyReport]:
    """Parse contractor stroyka PPTX → list of WeeklyReport."""
    if isinstance(path_or_bytes, (str, Path)):
        data = Path(path_or_bytes).read_bytes()
        filename = filename or Path(path_or_bytes).name
    else:
        data = path_or_bytes

    project_id = _project_id_from_name(filename)
    slide_titles = _slide_titles_from_pptx(data)

    reports: list[WeeklyReport] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        embeds = sorted(
            n for n in zf.namelist() if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")
        )
        for i, emb in enumerate(embeds, start=1):
            xbytes = zf.read(emb)
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                tmp.write(xbytes)
                tmp_path = tmp.name
            try:
                wb = load_workbook(tmp_path, data_only=True)
                ws = wb.active
                rep = _parse_sheet(ws, i)
            finally:
                Path(tmp_path).unlink(missing_ok=True)

            if i - 1 < len(slide_titles) and slide_titles[i - 1]:
                slide_title = slide_titles[i - 1]
                ole_low = rep.title_ole.lower()
                slide_low = slide_title.lower()
                # Prefer slide title when OLE is wrong (фахверк vs плита) or OLE empty
                use_slide = (
                    not rep.title_ole
                    or ("фахверк" in ole_low and "плит" in slide_low)
                    or (len(slide_title) > len(rep.title_ole) + 10 and "фундамент" in slide_low)
                )
                if use_slide:
                    if "фахверк" in ole_low and "плит" in slide_low:
                        rep.notes.append(
                            f"OLE-заголовок «{rep.title_ole}» не совпадает со слайдом — берём название слайда"
                        )
                    rep.title = slide_title
            rep.project_id = project_id
            # Fix typo борные → сборные for matching
            if "борные" in rep.title and "сборные" not in rep.title:
                rep.title = rep.title.replace("борные", "сборные")
                rep.notes.append("Исправлена опечатка «борные»→«сборные»")
            reports.append(rep)

    return reports
